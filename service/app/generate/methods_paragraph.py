from datetime import date
from io import BytesIO

from docx import Document
from docx.shared import Pt

from .models import GenerateRequest


def build_sentences(req: GenerateRequest) -> list[str]:
    sentences: list[str] = []

    opening = f"This {req.study_design} study" if req.study_design else "This study"
    if req.site:
        opening += f" will be conducted at {req.site}"
    if req.target_population:
        opening += f" among {req.target_population}"
    sentences.append(opening + ".")

    if req.inclusion_criteria:
        sentences.append(
            "Eligible participants must meet the following inclusion criteria: "
            + "; ".join(req.inclusion_criteria)
            + "."
        )

    if req.exclusion_criteria:
        sentences.append(
            "Participants will be excluded if they meet any of the following: "
            + "; ".join(req.exclusion_criteria)
            + "."
        )

    if req.calculated_sample_size:
        s = f"A total sample size of {req.calculated_sample_size} participants was determined necessary for this study"
        if req.sample_size_method:
            s += f", calculated based on {req.sample_size_method}"
        s += "."
        sentences.append(s)

    if req.primary_analysis:
        s = f"The primary analysis will be {req.primary_analysis}"
        if req.statistical_methods:
            s += f", using {req.statistical_methods}"
        if req.analysis_software:
            s += f" performed in {req.analysis_software}"
        if req.significance_level:
            s += f", with statistical significance set at alpha = {req.significance_level}"
        sentences.append(s + ".")

    return sentences


def build_methods_paragraph_docx(req: GenerateRequest) -> bytes:
    doc = Document()

    title = doc.add_heading(req.study_title, level=1)
    title.runs[0].font.size = Pt(16)
    doc.add_paragraph(f"Methods paragraph — Protocol v{req.version_number}")
    doc.add_paragraph(f"Generated {date.today().isoformat()}")
    doc.add_paragraph(
        "Drafted directly from the registered protocol fields. Review and adapt the "
        "wording before use — this is a starting point, not a finished manuscript section."
    )

    sentences = build_sentences(req)
    if not sentences:
        doc.add_paragraph("Not enough protocol fields are filled in yet to draft this paragraph.")
    else:
        doc.add_paragraph(" ".join(sentences))

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
