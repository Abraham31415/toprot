from datetime import date
from io import BytesIO

from docx import Document
from docx.shared import Pt

from .models import GenerateRequest


def build_consent_form_docx(req: GenerateRequest) -> bytes:
    doc = Document()

    title = doc.add_heading("Informed Consent Form", level=1)
    title.runs[0].font.size = Pt(16)
    doc.add_paragraph(req.study_title)
    doc.add_paragraph(f"Generated {date.today().isoformat()} — review and adapt before use.")

    doc.add_heading("Invitation", level=2)
    doc.add_paragraph(
        f"You are being invited to take part in a research study"
        + (f" conducted at {req.site}" if req.site else "")
        + (f" by {req.principal_investigator_name}" if req.principal_investigator_name else "")
        + ". Please read this form carefully before deciding whether to participate."
    )

    doc.add_heading("Purpose of the study", level=2)
    doc.add_paragraph(
        req.primary_analysis
        or "TODO: describe the purpose of this study in plain language."
    )

    doc.add_heading("Procedures", level=2)
    doc.add_paragraph(
        req.data_collection_methods
        or "TODO: describe what will happen if the participant agrees to take part."
    )

    doc.add_heading("Risks and benefits", level=2)
    doc.add_paragraph(req.risks_and_benefits or "TODO: describe risks, discomforts, and anticipated benefits.")

    doc.add_heading("Confidentiality", level=2)
    doc.add_paragraph(
        req.confidentiality_plan
        or "TODO: describe how participant data will be kept confidential."
    )

    doc.add_heading("Voluntary participation", level=2)
    doc.add_paragraph(
        "Taking part in this study is entirely voluntary. You may withdraw at any time "
        "without giving a reason and without any penalty or loss of benefits to which "
        "you are otherwise entitled."
    )

    doc.add_heading("Contact information", level=2)
    doc.add_paragraph(
        req.ethical_approval_body
        or "TODO: provide contact details for the research team and the ethics committee of record."
    )

    doc.add_heading("Statement of consent", level=2)
    doc.add_paragraph(
        "I have read and understood the information provided above. I have had the "
        "opportunity to ask questions, and I voluntarily agree to take part in this study."
    )
    doc.add_paragraph("Participant name: ______________________     Signature: ______________________     Date: ____________")
    doc.add_paragraph("Researcher name: ______________________     Signature: ______________________     Date: ____________")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
