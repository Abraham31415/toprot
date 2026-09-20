from datetime import date
from io import BytesIO

from docx import Document
from docx.shared import Pt

from .models import GenerateRequest, VariablePayload


def summary_row_label(v: VariablePayload) -> str:
    if v.variable_type == "continuous":
        return f"{v.variable_name}, mean (SD)"
    return f"{v.variable_name}, n (%)"


def build_table_shells_docx(req: GenerateRequest) -> bytes:
    doc = Document()

    title = doc.add_heading(req.study_title, level=1)
    title.runs[0].font.size = Pt(16)
    doc.add_paragraph(f"Table shells — Protocol v{req.version_number}")
    doc.add_paragraph(f"Generated {date.today().isoformat()}")
    doc.add_paragraph(
        "Empty shells for the tables planned in the statistical plan. Fill in once "
        "analysis is complete — structure, not content, is generated."
    )

    n_label = f"N = {req.calculated_sample_size}" if req.calculated_sample_size else "N = ___"

    covariates = [v for v in req.variables if v.role in ("covariate", "exposure", "identifier", "other")]
    outcomes = [v for v in req.variables if v.role == "outcome"]

    doc.add_heading(f"Table 1. Baseline characteristics ({n_label})", level=2)
    if not covariates:
        doc.add_paragraph("No covariate/exposure variables defined yet.")
    else:
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = "Characteristic", "Overall"
        for v in covariates:
            row = table.add_row().cells
            row[0].text = summary_row_label(v)
            row[1].text = ""

    doc.add_heading("Table 2. Primary analysis", level=2)
    if req.primary_analysis:
        doc.add_paragraph(req.primary_analysis)
    if not outcomes:
        doc.add_paragraph("No outcome variables defined yet.")
    else:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["Outcome", "Estimate", "95% CI", "p-value"]):
            hdr[i].text = h
        for v in outcomes:
            row = table.add_row().cells
            row[0].text = v.variable_name
            row[1].text = ""
            row[2].text = ""
            row[3].text = ""

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
