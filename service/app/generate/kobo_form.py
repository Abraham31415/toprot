from io import BytesIO

from openpyxl import Workbook

from .models import GenerateRequest

# XLSForm question type for each ToProt variable type. Anything with value
# labels becomes select_one regardless of this default (see build_kobo_form_xlsx).
DEFAULT_XLSFORM_TYPE = {
    "continuous": "decimal",
    "categorical": "text",
    "binary": "text",
    "ordinal": "text",
    "date": "date",
    "text": "text",
}


def build_kobo_form_xlsx(req: GenerateRequest) -> bytes:
    wb = Workbook()

    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])

    choices = wb.create_sheet("choices")
    choices.append(["list_name", "name", "label"])

    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id"])
    settings.append([req.study_title, f"toprot_v{req.version_number}"])

    for v in req.variables:
        if v.value_labels:
            list_name = f"{v.variable_slug}_choices"
            survey.append([f"select_one {list_name}", v.variable_slug, v.variable_name])
            for code, label in v.value_labels.items():
                choices.append([list_name, code, label])
        else:
            xls_type = DEFAULT_XLSFORM_TYPE.get(v.variable_type, "text")
            survey.append([xls_type, v.variable_slug, v.variable_name])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
