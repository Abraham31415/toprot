from datetime import date
from io import BytesIO
from typing import Optional

from docx import Document
from docx.shared import Pt

from .models import GenerateRequest

HEADERS = ["Variable", "Stata name", "Type", "Role", "Unit", "Value labels", "Timepoint", "Description"]


def format_value_labels(labels: Optional[dict[str, str]]) -> str:
    if not labels:
        return ""
    return "; ".join(f"{k} = {v}" for k, v in labels.items())


def build_data_dictionary_docx(req: GenerateRequest) -> bytes:
    doc = Document()

    title = doc.add_heading(req.study_title, level=1)
    title.runs[0].font.size = Pt(16)

    doc.add_paragraph(f"Data dictionary — Protocol v{req.version_number}")
    doc.add_paragraph(f"Generated {date.today().isoformat()}")

    if not req.variables:
        doc.add_paragraph("No variables have been defined yet.")
    else:
        table = doc.add_table(rows=1, cols=len(HEADERS))
        table.style = "Light Grid Accent 1"

        header_cells = table.rows[0].cells
        for i, header in enumerate(HEADERS):
            header_cells[i].text = header
            for p in header_cells[i].paragraphs:
                for run in p.runs:
                    run.font.bold = True

        for v in req.variables:
            row = table.add_row().cells
            row[0].text = v.variable_name
            row[1].text = v.variable_slug
            row[2].text = v.variable_type
            row[3].text = v.role
            row[4].text = v.unit or ""
            row[5].text = format_value_labels(v.value_labels)
            row[6].text = v.measurement_timepoint or ""
            row[7].text = v.description or ""

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
