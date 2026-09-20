from datetime import date
from io import BytesIO
from typing import Optional

from docx import Document
from docx.shared import Pt
from pydantic import BaseModel


class DeviationItem(BaseModel):
    category: str
    description: str
    occurred_at: str
    impact_assessment: Optional[str] = None


class AmendmentItem(BaseModel):
    reason: str
    created_at: str


class DeviationReportRequest(BaseModel):
    study_title: str
    version_number: int
    planned_sample_size: Optional[int] = None
    actual_enrollment: int
    study_design: Optional[str] = None
    primary_analysis: Optional[str] = None
    target_population: Optional[str] = None
    deviations: list[DeviationItem] = []
    amendments: list[AmendmentItem] = []


def add_bold(paragraph, text: str):
    run = paragraph.add_run(text)
    run.bold = True
    return run


def build_deviation_report_docx(req: DeviationReportRequest) -> bytes:
    doc = Document()

    title = doc.add_heading(req.study_title, level=1)
    title.runs[0].font.size = Pt(16)
    doc.add_paragraph(f"End-of-study deviation report — Registered protocol v{req.version_number}")
    doc.add_paragraph(f"Generated {date.today().isoformat()}")

    doc.add_heading("Registered protocol summary", level=2)
    if req.study_design:
        doc.add_paragraph(f"Design: {req.study_design}")
    if req.target_population:
        doc.add_paragraph(f"Target population: {req.target_population}")
    if req.primary_analysis:
        doc.add_paragraph(f"Primary analysis: {req.primary_analysis}")

    doc.add_heading("Sample size: planned vs. actual", level=2)
    table = doc.add_table(rows=2, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Planned", "Actual", "Difference"
    planned = req.planned_sample_size
    diff = (req.actual_enrollment - planned) if planned is not None else None
    row = table.rows[1].cells
    row[0].text = str(planned) if planned is not None else "Not recorded"
    row[1].text = str(req.actual_enrollment)
    row[2].text = str(diff) if diff is not None else "—"

    doc.add_heading("Logged deviations", level=2)
    if not req.deviations:
        doc.add_paragraph("No deviations were logged during the study.")
    else:
        dtable = doc.add_table(rows=1, cols=4)
        dtable.style = "Light Grid Accent 1"
        dhdr = dtable.rows[0].cells
        for i, h in enumerate(["Date", "Category", "Description", "Impact"]):
            dhdr[i].text = h
        for d in req.deviations:
            r = dtable.add_row().cells
            r[0].text = d.occurred_at
            r[1].text = d.category.replace("_", " ")
            r[2].text = d.description
            r[3].text = d.impact_assessment or ""

    doc.add_heading("Protocol amendments", level=2)
    if not req.amendments:
        doc.add_paragraph("The protocol was not amended after initial registration.")
    else:
        atable = doc.add_table(rows=1, cols=2)
        atable.style = "Light Grid Accent 1"
        ahdr = atable.rows[0].cells
        ahdr[0].text, ahdr[1].text = "Date", "Reason"
        for a in req.amendments:
            r = atable.add_row().cells
            r[0].text = a.created_at
            r[1].text = a.reason

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
