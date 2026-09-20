from datetime import date
from io import BytesIO
from typing import Optional

from docx import Document
from docx.shared import Pt

from .models import GenerateRequest


def add_field(doc: Document, label: str, value: Optional[str]):
    if value:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)


def build_ethics_package_docx(req: GenerateRequest) -> bytes:
    doc = Document()

    title = doc.add_heading(req.study_title, level=1)
    title.runs[0].font.size = Pt(16)
    doc.add_paragraph(f"Ethics submission summary — Protocol v{req.version_number}")
    doc.add_paragraph(f"Generated {date.today().isoformat()}")

    doc.add_heading("Study identity", level=2)
    add_field(doc, "Principal investigator", req.principal_investigator_name)
    add_field(doc, "Site", req.site)
    add_field(doc, "Study design", req.study_design)

    doc.add_heading("Population and eligibility", level=2)
    add_field(doc, "Target population", req.target_population)
    add_field(doc, "Setting", req.setting)
    if req.inclusion_criteria:
        doc.add_paragraph("Inclusion criteria:")
        for c in req.inclusion_criteria:
            doc.add_paragraph(c, style="List Bullet")
    if req.exclusion_criteria:
        doc.add_paragraph("Exclusion criteria:")
        for c in req.exclusion_criteria:
            doc.add_paragraph(c, style="List Bullet")

    doc.add_heading("Sample size", level=2)
    add_field(doc, "Calculated sample size", str(req.calculated_sample_size) if req.calculated_sample_size else None)
    add_field(doc, "Method", req.sample_size_method)

    doc.add_heading("Data collection and management", level=2)
    add_field(doc, "Data collection methods", req.data_collection_methods)
    add_field(doc, "Data management plan", req.data_management_plan)

    doc.add_heading("Ethical considerations", level=2)
    add_field(doc, "Ethical approval body", req.ethical_approval_body)
    add_field(doc, "Consent process", req.consent_process)
    add_field(doc, "Risks and benefits", req.risks_and_benefits)
    add_field(doc, "Confidentiality plan", req.confidentiality_plan)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
